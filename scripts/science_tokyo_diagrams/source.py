"""Native source creation and package-level deterministic normalization."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
import zipfile

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import DEFINITIONS, IDS
from .drawing import WIDTH, HEIGHT, WHITE, rgb


def canonical_zip(data):
    """Stable member ordering, ZIP timestamps, and nested chart workbooks."""
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for filename in sorted(src.namelist()):
            payload = src.read(filename)
            if filename.endswith(".xlsx"):
                payload = canonical_zip(payload)
            elif filename == "docProps/core.xml":
                payload = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                                 rb"\g<1>2000-01-01T00:00:00Z\g<2>", payload)
            member = zipfile.ZipInfo(filename, date_time=(2000, 1, 1, 0, 0, 0))
            member.compress_type = zipfile.ZIP_DEFLATED
            member.external_attr = 0o600 << 16
            dst.writestr(member, payload)
    return out.getvalue()


def generate_source(path):
    path = Path(path)
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(WIDTH), Inches(HEIGHT)
    props = ppt.core_properties
    props.title = "Science Tokyo editable diagram archetypes"
    props.subject = "Fictional academic planning examples; 22 native editable diagrams"
    props.author = "Science Tokyo template contributors"
    props.last_modified_by = props.author
    props.created = props.modified = datetime(2000, 1, 1)
    props.revision = 1
    for diagram_id, draw in DEFINITIONS:
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide._element.cSld.set("name", diagram_id)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(WHITE)
        draw(slide)
        slide.notes_slide.notes_text_frame.text = (
            f"diagram-id: {diagram_id}\n構成・数値は説明用の架空例。\n"
            "Yu Gothic。数値・図形・コネクタ・チャートは編集可能。")
    raw = BytesIO()
    ppt.save(raw)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(canonical_zip(raw.getvalue()))
    return validate_source(path)


def validate_source(path):
    ppt = Presentation(path)
    if len(ppt.slides) != len(IDS):
        raise ValueError(f"Source needs exactly 22 slides, got {len(ppt.slides)}")
    names = tuple(s._element.cSld.get("name") for s in ppt.slides)
    if names != IDS:
        raise ValueError("Source slide IDs/order do not match the 22-item image mapping")
    if abs(ppt.slide_width / 914400 - WIDTH) > .001 or abs(ppt.slide_height / 914400 - HEIGHT) > .001:
        raise ValueError("Source canvas must be 11.6 x 4.5 inches")
    counts = []
    for diagram_id, slide in zip(IDS, ppt.slides):
        if len(slide.shapes) < 5:
            raise ValueError(f"{diagram_id}: expected native editable shapes")
        if any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes):
            raise ValueError(f"{diagram_id}: flattened pictures are not editable diagram sources")
        counts.append({"id": diagram_id, "shapes": len(slide.shapes),
                       "charts": sum(s.has_chart for s in slide.shapes)})
    return counts
