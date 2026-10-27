"""Native chart builders with embedded workbooks and explicit linear scales."""
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                             XL_MARKER_STYLE, XL_TICK_MARK)
from pptx.util import Inches, Pt
from .drawing import NAVY, BLUE, RULE, PALE, INK, FONT, rgb, font_style

COLORS = (NAVY, BLUE, "B8C7E0")


def style(chart, *, ymin=0, ymax=100, step=20, legend=False):
    chart.has_title = False
    chart.has_legend = legend
    font_style(chart.font, 14)
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        font_style(chart.legend.font, 14)
    for axis in (chart.category_axis, chart.value_axis):
        font_style(axis.tick_labels.font, 14)
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = rgb(RULE)
    axis = chart.value_axis
    axis.minimum_scale, axis.maximum_scale, axis.major_unit = ymin, ymax, step
    axis.tick_labels.number_format = "0"
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = rgb(RULE)
    axis.major_gridlines.format.line.width = Pt(.6)
    for i, series in enumerate(chart.series):
        color = COLORS[i % len(COLORS)]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
        series.format.line.color.rgb = rgb(color)
        series.format.line.width = Pt(2.3)
    return chart


def category(slide, name, categories, series, x, y, w, h, *, kind="column",
             ymax=100, step=20, labels=True, legend=False):
    data = CategoryChartData()
    data.categories = categories
    for label, values in series:
        data.add_series(label, values)
    kinds = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED,
             "bar": XL_CHART_TYPE.BAR_CLUSTERED,
             "stacked": XL_CHART_TYPE.COLUMN_STACKED,
             "line": XL_CHART_TYPE.LINE_MARKERS}
    frame = slide.shapes.add_chart(kinds[kind], Inches(x), Inches(y),
                                  Inches(w), Inches(h), data)
    frame.name = f"chart:{name}"
    chart = style(frame.chart, ymax=ymax, step=step, legend=legend)
    plot = chart.plots[0]
    if kind in ("column", "bar", "stacked"):
        plot.gap_width = 80
    if kind == "stacked":
        plot.overlap = 100
    if kind == "bar":
        chart.category_axis.reverse_order = True
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        font_style(dl.font, 14, color="FFFFFF" if kind == "stacked" else INK)
        dl.position = (XL_LABEL_POSITION.CENTER if kind == "stacked"
                       else XL_LABEL_POSITION.OUTSIDE_END)
        dl.number_format = "0"
        if kind == "stacked":
            for i, series in enumerate(chart.series):
                for point in series.points:
                    font_style(point.data_label.font, 14, color="FFFFFF" if i == 0 else INK)
    if kind == "line":
        for series in chart.series:
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 7
    return chart


def xy(slide, name, series, x, y, w, h, *, xmin, xmax, xstep,
       ymin, ymax, ystep, lines=False, legend=True):
    data = XyChartData()
    for label, points in series:
        s = data.add_series(label)
        for px, py in points:
            s.add_data_point(px, py)
    kind = XL_CHART_TYPE.XY_SCATTER_LINES if lines else XL_CHART_TYPE.XY_SCATTER
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = f"chart:{name}"
    chart = style(frame.chart, ymin=ymin, ymax=ymax, step=ystep, legend=legend)
    axis = chart.category_axis
    axis.minimum_scale, axis.maximum_scale, axis.major_unit = xmin, xmax, xstep
    axis.tick_labels.number_format = "0"
    for i, s in enumerate(chart.series):
        s.marker.style = (XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE,
                          XL_MARKER_STYLE.DIAMOND)[i % 3]
        s.marker.size = 9 if not lines else 6
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = rgb(COLORS[i % 3])
        s.marker.format.line.color.rgb = rgb(COLORS[i % 3])
        if not lines:
            s.format.line.fill.background()
    return chart
