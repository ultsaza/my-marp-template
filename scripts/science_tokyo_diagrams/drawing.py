"""Native DrawingML primitives in inches; no pictures or rendering shortcuts."""
from __future__ import annotations
from math import sqrt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

NAVY = "1C3077"
BLUE = "7F96C2"
PALE = "E9EEF7"
PALE2 = "F4F6FA"
INK = "22304D"
MUTED = "52617D"
RULE = "C9D3E5"
WHITE = "FFFFFF"
FONT = "Yu Gothic"
WIDTH, HEIGHT = 11.6, 4.5


def rgb(value):
    return RGBColor.from_string(value)


def font_style(font, size=18, bold=False, color=INK):
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = rgb(color)


def text(slide, value, x, y, w, h=.42, *, size=18, bold=False,
         color=INK, align=PP_ALIGN.LEFT, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or f"text:{value[:32]}"
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.015)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.13
        run = p.add_run()
        run.text = line
        font_style(run.font, size, bold, color)
        # Preserve the existing family for CJK as well as Latin runs.
        ea = OxmlElement("a:ea")
        ea.set("typeface", FONT)
        run._r.get_or_add_rPr().append(ea)
    return shape


def disable_effects(item):
    # python-pptx's default shape style references theme effect 2 (a shadow).
    # An explicit empty effect list plus a null theme reference avoids that
    # inherited decoration in PowerPoint AND LibreOffice.
    for ref in item._element.xpath("./p:style/a:effectRef"):
        ref.set("idx", "0")
    sp_pr = item._element.spPr
    for old in list(sp_pr):
        if old.tag.endswith("}effectLst") or old.tag.endswith("}effectDag"):
            sp_pr.remove(old)
    sp_pr.append(OxmlElement("a:effectLst"))
    return item


def shape(slide, x, y, w, h, *, fill=PALE, line=None,
          kind=MSO_SHAPE.RECTANGLE, name="shape", line_width=1):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.name = name
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = rgb(fill)
    else:
        item.fill.background()
    if line:
        item.line.color.rgb = rgb(line)
        item.line.width = Pt(line_width)
    else:
        item.line.fill.background()
    return disable_effects(item)


def box(slide, value, x, y, w, h=.65, *, fill=PALE, color=INK,
        size=18, bold=False, line=None, name=None):
    base = shape(slide, x, y, w, h, fill=fill, line=line,
                 name=name or f"box:{value[:24]}")
    text(slide, value, x+.09, y+.035, w-.18, h-.07, size=size,
         bold=bold, color=color, align=PP_ALIGN.CENTER)
    return base


def line(slide, x1, y1, x2, y2, *, color=RULE, width=1.1,
         arrow=False, name="connector", dashed=False):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    item.name = name
    item.line.color.rgb = rgb(color)
    item.line.width = Pt(width)
    ln = item._element.spPr.get_or_add_ln()
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "med")
        end.set("len", "med")
        ln.append(end)
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        ln.append(dash)
    return disable_effects(item)


def arrow(slide, x1, y1, x2, y2, **kwargs):
    return line(slide, x1, y1, x2, y2, color=NAVY, width=1.8, arrow=True, **kwargs)


def circle(slide, cx, cy, diameter, *, fill=NAVY, line_color=None, name="circle"):
    return shape(slide, cx-diameter/2, cy-diameter/2, diameter, diameter,
                 fill=fill, line=line_color, kind=MSO_SHAPE.OVAL, name=name)


def area_diameter(value, maximum, max_diameter):
    """Circle AREA, not radius, encodes a non-negative value."""
    if value < 0 or maximum <= 0 or value > maximum:
        raise ValueError("Area value must be in 0..maximum")
    return max_diameter * sqrt(value/maximum)


def caption(slide, value):
    text(slide, value, .3, .08, 11, .4, size=17, bold=True, color=NAVY)


def key(slide, value, x, y, *, fill=NAVY, w=2.4):
    shape(slide, x, y+.13, .17, .17, fill=fill, name=f"legend:{value}")
    text(slide, value, x+.25, y, w-.25, .43, size=14, color=MUTED)


def axes(slide, x, y, w, h, ymax, step, *, unit, ticks=True):
    """Linear, zero-origin chart grid; returns a value-to-y function."""
    text(slide, unit, x-.1, y-.49, w, .4, size=14, color=MUTED)
    for v in range(0, ymax+1, step):
        yy = y+h-h*v/ymax
        line(slide, x, yy, x+w, yy, color=RULE if v else MUTED)
        if ticks:
            text(slide, str(v), x-.65, yy-.18, .5, .36, size=14,
                 align=PP_ALIGN.RIGHT, color=MUTED)
    line(slide, x, y, x, y+h, color=MUTED)
    return lambda value: y+h-h*value/ymax


def diamond(slide, value, x, y, w=2.2, h=1.25):
    shape(slide, x, y, w, h, fill=PALE, line=NAVY,
          kind=MSO_SHAPE.DIAMOND, name=f"decision:{value}")
    text(slide, value, x+w*.15, y+h*.24, w*.7, h*.52,
         size=18, bold=True, align=PP_ALIGN.CENTER)
