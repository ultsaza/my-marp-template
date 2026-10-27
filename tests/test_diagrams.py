"""The diagram source is editable; the Marp mapping is a stable 22-item API."""
from __future__ import annotations
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "scripts/build-science-tokyo-diagrams.py"
IDS = (
    "dot_matrix_share", "progress_bubble_matrix", "ranked_bar_annotated",
    "scatter_annotated", "proportional_circles", "delta_bars_totals",
    "scenario_lines_cagr", "swimlane_numbered_notes", "front_back_two_tier",
    "chart_insight", "stacked_bar", "waterfall", "true_waterfall",
    "small_multiples", "matrix_2x2", "issue_to_solution_map", "issue_tree",
    "calc_flow", "cycle", "decision_fork", "roadmap", "gantt",
)
SOURCE = "science-tokyo-diagrams.pptx"
SCRATCH = ROOT / ".verification/diagrams/tests"


def invoke(*args, check=True):
    assert CLI.is_file(), "Missing editable diagram generation/export CLI"
    return subprocess.run([sys.executable, str(CLI), *map(str, args)],
                          cwd=ROOT, capture_output=True, text=True, check=check)


class DiagramSourceTests(unittest.TestCase):
    def workspace(self):
        SCRATCH.mkdir(parents=True, exist_ok=True)
        folder = tempfile.TemporaryDirectory(dir=SCRATCH)
        self.addCleanup(folder.cleanup)
        return Path(folder.name)

    def test_source_has_exactly_22_ordered_native_editable_slides(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        ppt = Presentation(out / SOURCE)
        self.assertEqual(len(ppt.slides), 22)
        self.assertEqual((ppt.slide_width, ppt.slide_height), (Inches(11.6), Inches(4.5)))
        self.assertEqual([s._element.cSld.get("name") for s in ppt.slides], list(IDS))
        for diagram_id, slide in zip(IDS, ppt.slides):
            with self.subTest(diagram_id=diagram_id):
                self.assertGreaterEqual(len(slide.shapes), 5)
                self.assertFalse(any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes))
                self.assertTrue(any(s.has_text_frame and s.text.strip() for s in slide.shapes))
                self.assertTrue(all(s.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX,
                                                    MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.CHART)
                                    for s in slide.shapes))

    def test_stable_png_mapping_uses_contract_ids_and_slide_order(self):
        mapping = json.loads(invoke("--list").stdout)
        self.assertEqual(mapping, [{"id": name, "slide": i, "file": f"{name}.png"}
                                   for i, name in enumerate(IDS, 1)])

    def test_editable_shapes_stay_inside_canvas_and_text_is_readable(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        ppt = Presentation(out / SOURCE)
        for slide in ppt.slides:
            for sh in slide.shapes:
                label = f"{slide._element.cSld.get('name')}: {sh.name}"
                self.assertGreaterEqual(sh.left, 0, label)
                self.assertGreaterEqual(sh.top, 0, label)
                self.assertLessEqual(sh.left+sh.width, ppt.slide_width, label)
                self.assertLessEqual(sh.top+sh.height, ppt.slide_height, label)
                self.assertFalse(sh._element.xpath('.//a:prstGeom[@prst="roundRect"]'), label)
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                self.assertEqual(run.font.name, "Yu Gothic", label)
                                self.assertGreaterEqual(run.font.size.pt, 14, label)

    def test_native_shapes_disable_inherited_theme_shadows(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        for slide in Presentation(out / SOURCE).slides:
            for sh in slide.shapes:
                for ref in sh._element.xpath("./p:style/a:effectRef"):
                    self.assertEqual(ref.get("idx"), "0", f"Inherited shadow on {sh.name}")
                self.assertFalse(sh._element.xpath(".//a:outerShdw | .//a:innerShdw"))

    def test_negative_decision_routes_terminate_in_distinct_outcomes(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        slide = Presentation(out / SOURCE).slides[19]
        routes = {sh.name: sh for sh in slide.shapes if sh.name.startswith("decision-no:")}
        self.assertEqual(set(routes), {"decision-no:consent", "decision-no:anonymize"})
        for key, label in (("consent", "同意を確認"), ("anonymize", "匿名化を実施")):
            route = routes[f"decision-no:{key}"]
            outcome = next(sh for sh in slide.shapes if sh.name == f"box:{label}")
            self.assertEqual(route.begin_x, route.end_x)
            self.assertTrue(outcome.left < route.end_x < outcome.left+outcome.width)
            self.assertLessEqual(abs(route.end_y-outcome.top), Inches(.1))

    def test_gantt_grid_and_dependency_arrows_are_not_hidden_by_row_surfaces(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        slide = Presentation(out / SOURCE).slides[21]
        names = [sh.name for sh in slide.shapes]
        rows = [i for i, name in enumerate(names) if name.startswith("gantt-row:")]
        grid = [i for i, name in enumerate(names) if name.startswith("gantt-grid:")]
        deps = [i for i, name in enumerate(names) if name.startswith("dependency:")]
        self.assertEqual((len(rows), len(grid), len(deps)), (4, 7, 3))
        self.assertLess(max(rows), min(grid))
        self.assertLess(max(rows), min(deps))

    def test_stacked_labels_have_contrast_on_pale_segments(self):
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        chart = next(sh.chart for sh in Presentation(out / SOURCE).slides[10].shapes if sh.has_chart)
        for i, series in enumerate(chart.series):
            for point in series.points:
                color = point.data_label.font.color
                self.assertIsNotNone(color.type, "Each stack label needs explicit contrast, not global white")
                self.assertEqual(str(color.rgb), "FFFFFF" if i == 0 else "22304D")

    def test_quantitative_geometry_and_native_workbook_data(self):
        import zipfile
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        ppt = Presentation(out / SOURCE)
        with zipfile.ZipFile(out / SOURCE) as package:
            self.assertEqual(len([n for n in package.namelist() if n.endswith(".xlsx")]), 5)
            self.assertFalse(any(n.startswith("ppt/media/") for n in package.namelist()))
        charts = {sh.name: sh.chart for slide in ppt.slides for sh in slide.shapes if sh.has_chart}
        self.assertEqual(len(charts), 5)
        for chart in charts.values():
            self.assertEqual(chart.value_axis.minimum_scale, 0)
        self.assertEqual(tuple(charts["chart:applications"].series[0].values), (18, 14, 10, 6))
        scenario = charts["chart:scenarios"]
        for series, rate in zip(scenario.series, (.20, .10, .05)):
            xs = [float(n.text) for n in series._element.xpath("./c:xVal/c:numRef/c:numCache/c:pt/c:v")]
            self.assertEqual(xs, [2026, 2027, 2028, 2029])
            expected = [100*(1+rate)**i for i in range(4)]
            for actual, value in zip(series.values, expected):
                self.assertAlmostEqual(actual, value)
            self.assertAlmostEqual((series.values[-1]/series.values[0])**(1/3)-1, rate)
        circles = [sh for sh in ppt.slides[4].shapes if sh.name.startswith("area:")]
        self.assertEqual(len(circles), 3)
        base_area = circles[0].width*circles[0].height
        for sh, ratio in zip(circles, (1, 2, 4)):
            self.assertAlmostEqual(sh.width*sh.height/base_area, ratio, places=4)
        for index, maximum, expected in ((11, 80, ((0,40),(40,60),(60,50),(50,65),(0,65))),
                                         (12, 120, ((0,80),(80,110),(110,90),(90,100),(0,100)))):
            bars = [sh for sh in ppt.slides[index].shapes if sh.name.startswith("bridge:")]
            self.assertEqual(len(bars), 5)
            for sh, (before, after) in zip(bars, expected):
                self.assertEqual(tuple(map(int, sh.name.split(":")[-2:])), (before, after))
                top = 1.05+2.62*(1-max(before, after)/maximum)
                height = 2.62*abs(after-before)/maximum
                self.assertAlmostEqual(sh.top/Inches(1), top, places=5)
                self.assertAlmostEqual(sh.height/Inches(1), height, places=5)
        points = [sh for sh in ppt.slides[13].shapes if sh.name.startswith("point:")]
        self.assertEqual(len(points), 12)
        for sh in points:
            value = int(sh.name.rsplit(":", 1)[1])
            self.assertAlmostEqual((sh.top+sh.height/2)/Inches(1),
                                   1.56+1.97*(1-value/100), places=5)
        for sh in ppt.slides[21].shapes:
            if sh.name.startswith("gantt:"):
                start, duration = map(int, sh.name.split(":")[-2:])
                self.assertAlmostEqual(sh.left/Inches(1), 2.95+start*1.3, places=5)
                self.assertAlmostEqual(sh.width/Inches(1), duration*1.3, places=5)
        dots = [sh for sh in ppt.slides[0].shapes if sh.name.startswith("record:")]
        self.assertEqual(len(dots), 100)
        self.assertEqual(sum(sh.name.endswith(":open") for sh in dots), 64)

    def test_export_failure_is_nonzero_and_does_not_publish_assets(self):
        import os
        out = self.workspace()
        invoke("--generate-only", "--output-dir", out)
        fake_bin = out / "bin"
        fake_bin.mkdir()
        (fake_bin / "soffice").symlink_to("/bin/false")
        for tool in ("pdfinfo", "pdftoppm"):
            (fake_bin / tool).symlink_to(f"/usr/bin/{tool}")
        failed = subprocess.run([sys.executable, str(CLI), "--export-only", "--output-dir", str(out)],
                                capture_output=True, text=True, env=dict(os.environ, PATH=str(fake_bin)))
        self.assertNotEqual(failed.returncode, 0)
        self.assertIn("soffice", failed.stderr)
        self.assertFalse((out / "manifest.json").exists())
        self.assertEqual(list(out.glob("*.png")), [])

    def test_check_rejects_missing_and_corrupted_png_and_wrong_source_ids(self):
        import shutil
        committed = ROOT / "assets/science-tokyo/diagrams"
        out = self.workspace()
        for p in committed.iterdir():
            if p.is_file():
                shutil.copyfile(p, out / p.name)
        png = out / f"{IDS[0]}.png"
        png.unlink()
        missing = invoke("--check", "--output-dir", out, check=False)
        self.assertNotEqual(missing.returncode, 0)
        self.assertIn("missing", missing.stderr)
        png.write_bytes(b"not a valid PNG")
        corrupt = invoke("--check", "--output-dir", out, check=False)
        self.assertNotEqual(corrupt.returncode, 0)
        self.assertIn("image", corrupt.stderr.lower())
        ppt = Presentation(out / SOURCE)
        ppt.slides[0]._element.cSld.set("name", "unexpected")
        ppt.save(out / SOURCE)
        wrong = invoke("--export-only", "--output-dir", out, check=False)
        self.assertNotEqual(wrong.returncode, 0)
        self.assertIn("IDs/order", wrong.stderr)

    def test_generation_is_byte_deterministic(self):
        import time
        first, second = self.workspace(), self.workspace()
        invoke("--generate-only", "--output-dir", first)
        time.sleep(1.1)  # Catch workbook core timestamps, not merely ZIP dates.
        invoke("--generate-only", "--output-dir", second)
        self.assertEqual((first / SOURCE).read_bytes(), (second / SOURCE).read_bytes(),
                         "Native source generation must be byte-deterministic, including embedded XLSX")

    def test_repository_assets_are_in_sync(self):
        folder = ROOT / "assets/science-tokyo/diagrams"
        self.assertTrue((folder / SOURCE).is_file(), "Missing deliverable native PPTX source")
        checked = invoke("--check", "--output-dir", folder, check=False)
        self.assertEqual(checked.returncode, 0, checked.stderr)
        report = json.loads(checked.stdout)
        self.assertEqual((report["slides"], report["images"]), (22, 22))

    def test_default_build_renders_22_and_export_only_preserves_manual_edit(self):
        import hashlib
        from PIL import Image, ImageChops, ImageStat
        out = self.workspace()
        built = invoke("--output-dir", out, check=False)
        self.assertEqual(built.returncode, 0, built.stderr)
        result = json.loads(built.stdout)
        self.assertEqual((result["slides"], result["images"]), (22, 22))
        self.assertEqual({p.stem for p in out.glob("*.png")}, set(IDS))
        manifest = json.loads((out / "manifest.json").read_text())
        self.assertEqual([p["id"] for p in manifest["images"]], list(IDS))
        for item in manifest["images"]:
            with Image.open(out / item["file"]) as im:
                self.assertGreaterEqual(im.width, 1600)
                self.assertLess(abs(im.height/im.width - 4.5/11.6), .001)
                self.assertEqual(im.size, (item["width"], item["height"]))
                self.assertGreater(ImageStat.Stat(im.convert("L")).stddev[0], 15)
                ink = ImageChops.difference(im.convert("RGB"), Image.new("RGB", im.size, "white"))
                bbox = ink.getbbox()
                self.assertIsNotNone(bbox)
                self.assertGreaterEqual(bbox[0], 8)
                self.assertGreaterEqual(bbox[1], 8)
                self.assertLess(bbox[2], im.width-8)
                self.assertLess(bbox[3], im.height-8)
        self.assertEqual(invoke("--check", "--output-dir", out).returncode, 0)
        first_hashes = [row["sha256"] for row in manifest["images"]]
        invoke("--export-only", "--output-dir", out)
        repeated = json.loads((out / "manifest.json").read_text())
        self.assertEqual([row["sha256"] for row in repeated["images"]], first_hashes,
                         "Repeated exports on this renderer/font environment must be byte-identical")
        ppt = Presentation(out / SOURCE)
        target = next(sh for sh in ppt.slides[0].shapes if sh.has_text_frame and "64%" in sh.text)
        target.text = "MANUAL 64%"
        ppt.save(out / SOURCE)
        edited = (out / SOURCE).read_bytes()
        stale = invoke("--check", "--output-dir", out, check=False)
        self.assertNotEqual(stale.returncode, 0)
        self.assertIn("source", stale.stderr.lower())
        invoke("--export-only", "--output-dir", out)
        self.assertEqual((out / SOURCE).read_bytes(), edited)
        self.assertEqual(json.loads((out / "manifest.json").read_text())["source"]["sha256"],
                         hashlib.sha256(edited).hexdigest())
        self.assertEqual(invoke("--check", "--output-dir", out).returncode, 0)


if __name__ == "__main__":
    unittest.main()
